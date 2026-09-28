# Turns the hand-over written by build-templates.mjs into the Excel template.
# Usage: python build-templates.py <handover.json> <output.xlsx>
import datetime as dt
import json
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

with open(sys.argv[1], encoding='utf-8') as f:
    data = json.load(f)

HEAD_FONT = Font(bold=True, color='FFFFFF')
HEAD_FILL = PatternFill('solid', fgColor='131A27')


def sheet(wb, title, rows, date_col=None, widths=None, first=False):
    ws = wb.active if first else wb.create_sheet(title)
    ws.title = title
    for i, r in enumerate(rows):
        row = list(r)
        if i > 0 and date_col is not None:
            row[date_col] = dt.date.fromisoformat(row[date_col])
        ws.append(row)
    for c in range(1, len(rows[0]) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = HEAD_FONT
        cell.fill = HEAD_FILL
    if date_col is not None:
        for row in ws.iter_rows(min_row=2, min_col=date_col + 1, max_col=date_col + 1):
            row[0].number_format = 'YYYY-MM-DD'
    for i, w in enumerate(widths or [], start=1):
        ws.column_dimensions[chr(64 + i)].width = w
    ws.freeze_panes = 'A2'
    return ws


wb = Workbook()
sheet(wb, 'Transactions', data['transactions'], date_col=0, widths=[13, 12, 12, 14, 11], first=True)
sheet(wb, 'Opening stock', data['opening'], date_col=2, widths=[12, 12, 13, 11])
sheet(wb, 'Products', data['products'], widths=[12, 24, 12, 16, 11, 11, 15, 11])

info = wb.create_sheet('Read me')
info.column_dimensions['A'].width = 30
info.column_dimensions['B'].width = 112
TEXT = [
    ('What this file is', 'A template for the Simon G. Stock Check. Replace the example rows with your own and load the file on groma.ro/en/simon-g/stock-check. The file is read on your computer and is not uploaded.'),
    ('Ce este acest fișier', 'Un model pentru verificarea de stoc Simon G. Înlocuiți rândurile de exemplu cu ale dumneavoastră și încărcați fișierul pe groma.ro/simon-g/verificare-stoc. Fișierul este citit pe calculatorul dumneavoastră și nu este trimis nicăieri.'),
    ('Transactions', 'One row per movement: date, product, location, type, quantity. Types: sale, receipt (from the supplier), return (from a customer), adjustment (plus or minus). Quantities are positive; the type gives the direction. Only an adjustment carries a sign.'),
    ('Tranzacții', 'Câte un rând pe mișcare: data, produsul, locația, tipul, cantitatea. Tipuri: sale (vânzare), receipt (recepție de la furnizor), return (retur de la client), adjustment (ajustare, cu plus sau minus). Cantitățile sunt pozitive; direcția o dă tipul. Doar ajustarea are semn.'),
    ('Opening stock', 'One row per product and location: the stock at the start, with its date. From it and from the transactions the page rebuilds the stock of every day, so it knows when a product ran out.'),
    ('Stoc inițial', 'Câte un rând pe produs și locație: stocul de la început, cu data lui. Din el și din tranzacții pagina reface stocul fiecărei zile, așa că știe când s-a terminat un produs.'),
    ('Products (optional)', 'One row per product: lead time in days, unit cost, pack size, minimum order quantity, quantity already on order. Any column may be left empty; the page then uses the assumption you set on screen and says so.'),
    ('Produse (opțional)', 'Câte un rând pe produs: termenul de livrare în zile, costul unitar, bucățile pe ambalaj, cantitatea minimă de comandă, cantitatea deja comandată. Orice coloană poate rămâne goală; pagina folosește atunci ipoteza aleasă pe ecran și spune acest lucru.'),
    ('Your own column names', 'You may keep the column names of your system. The page shows which column it took for what and lets you change it. It also lets you say what each of your transaction types means.'),
    ('Numele coloanelor', 'Puteți păstra numele de coloane din sistemul dumneavoastră. Pagina arată ce coloană a luat pentru ce și vă lasă să schimbați. Tot acolo spuneți ce înseamnă fiecare tip de tranzacție de la dumneavoastră.'),
    ('How much history', 'At least three months. One to two years is better; two full years allow a yearly pattern to be measured.'),
    ('Cât istoric', 'Cel puțin trei luni. Unul-doi ani este mai bine; doi ani întregi permit măsurarea unui tipar anual.'),
    ('The example data', data.get('note_en', 'Invented: 24 products of a parts distributor, one location, one year.')),
    ('Datele de exemplu', data.get('note_ro', 'Inventate: 24 de produse ale unui distribuitor de piese, o locație, un an.')),
]
for k, (title, text) in enumerate(TEXT, start=1):
    info.cell(row=k, column=1, value=title).font = Font(bold=True)
    info.cell(row=k, column=1).alignment = Alignment(vertical='top')
    c = info.cell(row=k, column=2, value=text)
    c.alignment = Alignment(wrap_text=True, vertical='top')

wb.save(sys.argv[2])
